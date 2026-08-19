from __future__ import annotations

from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from routes.context import quotation_repo, window_snapshot_service
from services.notes import load_notes


HEADERS = [
    "Name",
    "Role",
    "Team",
    "QI",
    "QA",
    "Delta",
    "FVM",
    "W QI",
    "W DELTA",
    "W FV",
    "W MV",
    "W Goals",
    "W Assists",
    "W Yellow",
    "W Red",
    "W Records",
    "Notes",
]


def build_dashboard_export(
    *,
    season: int | None,
    days: int = 38,
    role: str | None = None,
    team: str | None = None,
    search: str | None = None,
    sort: str = "FVM",
) -> Workbook:
    if season is None:
        season = quotation_repo.find_latest_season()
    if season is None:
        raise RuntimeError("No season available")

    payload = window_snapshot_service.list_rows(
        season,
        days,
        role=role,
        team=team,
        search=search,
        sort=sort,
        direction=-1,
        limit=10000,
    )

    rows = [_export_row(row, load_notes()) for row in payload["rows"]]

    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    _write_sheet(ws, rows)

    targets = sorted(
        rows,
        key=lambda row: _num(row[8]),
        reverse=True,
    )
    ws_targets = wb.create_sheet("Auction Targets")
    _write_sheet(ws_targets, targets)

    undervalued = [row for row in rows if _num(row[8]) > 0]
    ws_under = wb.create_sheet("Positive W DELTA")
    _write_sheet(ws_under, undervalued)

    watchlist = [row for row in rows if str(row[-1]).strip()]
    ws_watch = wb.create_sheet("Watchlist")
    _write_sheet(ws_watch, watchlist)

    for sheet in wb.worksheets:
        _style_sheet(sheet)

    return wb


def _export_row(row: dict[str, Any], notes: dict[str, str]) -> list:
    player_id = row.get("player_id")
    name = row.get("name")
    qa = _num(row.get("QA"))
    w_qi = _num(row.get("window_QI"))
    w_delta = None if qa is None or w_qi is None else w_qi - qa
    note = notes.get(str(player_id), "") or notes.get(str(name), "")

    return [
        name,
        row.get("role"),
        row.get("team"),
        row.get("QI"),
        row.get("QA"),
        row.get("quotation_delta"),
        row.get("FVM"),
        row.get("window_QI"),
        w_delta,
        row.get("window_FV"),
        row.get("window_MV"),
        row.get("window_goals"),
        row.get("window_assists"),
        row.get("window_yellow_cards"),
        row.get("window_red_cards"),
        row.get("window_records"),
        note,
    ]


def _write_sheet(ws, rows: list[list[Any]]) -> None:
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    even_fill = PatternFill("solid", fgColor="F3F6FA")
    thin = Side(style="thin", color="D9E2EF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Header row only.
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    # Body rows.
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

            if row_idx % 2 == 0:
                cell.fill = even_fill

    # W DELTA column = I.
    if ws.max_row >= 2:
        ws.conditional_formatting.add(
            f"I2:I{ws.max_row}",
            CellIsRule(
                operator="greaterThan",
                formula=["0"],
                fill=PatternFill("solid", fgColor="C6EFCE"),
                font=Font(color="006100"),
            ),
        )
        ws.conditional_formatting.add(
            f"I2:I{ws.max_row}",
            CellIsRule(
                operator="lessThan",
                formula=["0"],
                fill=PatternFill("solid", fgColor="FFC7CE"),
                font=Font(color="9C0006"),
            ),
        )

    for column in ws.columns:
        max_len = 0
        letter = get_column_letter(column[0].column)

        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 60))

        ws.column_dimensions[letter].width = max(10, max_len + 2)


def _num(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return -10**12
